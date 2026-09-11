"""
Consolidated employee workbook (.xlsx) -> data.json

Trigger:  EventBridge "Object Created" on the consolidated/ prefix of an S3 bucket.
Action:   read the workbook, map the consolidated table headers onto the target
          JSON keys, and write data.json back to S3.

Output (OUTPUT_FORMAT=js_const, the default):

    const EMPLOYEES = [
      {
        "login": "jdelacruz",
        "name": "Juan Dela Cruz",
        "email": "juandelacruz@example.com",
        "manager": "Raymart Dela Cruz",
        "market": "HKIT",
        "marketTag": "HKIT",
        "active": true,
        "tokens": 0.0
      }
    ];

Set OUTPUT_FORMAT=json to emit a plain JSON array instead.

Runtime: python3.12+. Needs openpyxl (AWS SDK for pandas managed layer includes it,
or ship your own layer -- see README.md).
"""

from __future__ import annotations

import datetime as dt
import io
import json
import logging
import os
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Iterable, Iterator
from urllib.parse import unquote_plus

import boto3
from botocore.exceptions import ClientError
from openpyxl import load_workbook

LOGGER = logging.getLogger()
LOGGER.setLevel(os.environ.get("LOG_LEVEL", "INFO"))

s3 = boto3.client("s3")


# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

def _env(name: str, default: str = "") -> str:
    return os.environ.get(name, default).strip()


def _env_bool(name: str, default: bool = False) -> bool:
    raw = _env(name)
    if not raw:
        return default
    return raw.lower() in {"1", "true", "yes", "on"}


def _env_int(name: str, default: int) -> int:
    raw = _env(name)
    try:
        return int(raw) if raw else default
    except ValueError:
        LOGGER.warning("Env %s=%r is not an integer; using %d", name, raw, default)
        return default


def _env_json(name: str, default: Any) -> Any:
    raw = _env(name)
    if not raw:
        return default
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        LOGGER.error("Env %s is not valid JSON; using default", name)
        return default


@dataclass(frozen=True)
class Config:
    # --- where things live -------------------------------------------------
    source_prefix: str = _env("SOURCE_PREFIX", "consolidated/")
    output_bucket: str = _env("OUTPUT_BUCKET")          # blank -> same bucket as input
    output_key: str = _env("OUTPUT_KEY", "output/data.json")
    report_key: str = _env("REPORT_KEY", "output/last-run-report.json")
    sheet_name: str = _env("SHEET_NAME")                # blank -> first/active sheet

    # --- output shape ------------------------------------------------------
    output_format: str = _env("OUTPUT_FORMAT", "js_const").lower()   # js_const | json
    js_const_name: str = _env("JS_CONST_NAME", "EMPLOYEES")
    sort_by_login: bool = _env_bool("SORT_BY_LOGIN", True)

    # --- field derivation --------------------------------------------------
    # active: last_active_window | always_true | employee_type
    active_mode: str = _env("ACTIVE_MODE", "last_active_window").lower()
    active_window_days: int = _env_int("ACTIVE_WINDOW_DAYS", 90)
    active_blank_date: bool = _env_bool("ACTIVE_IF_DATE_BLANK", False)
    inactive_employee_types: tuple = tuple(
        t.strip().lower()
        for t in _env("INACTIVE_EMPLOYEE_TYPES", "terminated,leaver,inactive").split(",")
        if t.strip()
    )

    # market: which consolidated column feeds it, and how values translate
    market_source_column: str = _env("MARKET_SOURCE_COLUMN", "business category")
    market_map: dict = field(default_factory=lambda: _env_json("MARKET_MAP", {}))
    market_default: str = _env("MARKET_DEFAULT", "")
    market_strict: bool = _env_bool("MARKET_STRICT", False)
    market_uppercase: bool = _env_bool("MARKET_UPPERCASE", True)

    default_tokens: float = float(_env("DEFAULT_TOKENS", "0.0"))
    derive_login_from_email: bool = _env_bool("DERIVE_LOGIN_FROM_EMAIL", False)

    # --- safety ------------------------------------------------------------
    require_email: bool = _env_bool("REQUIRE_EMAIL", True)
    max_reject_ratio: float = float(_env("MAX_REJECT_RATIO", "0.25"))
    sse: str = _env("S3_SSE", "AES256")                 # AES256 | aws:kms | "" to disable
    sse_kms_key_id: str = _env("S3_SSE_KMS_KEY_ID")


CFG = Config()

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


# --------------------------------------------------------------------------- #
# Header handling
# --------------------------------------------------------------------------- #

def canon(value: Any) -> str:
    """Canonical form of a header cell: 'Business  Category ' -> 'business category'."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace(" ", " ")
    text = re.sub(r"[^0-9a-zA-Z]+", " ", text)
    return re.sub(r"\s+", " ", text).strip().lower()


# Consolidated header -> internal name. Aliases let the BA's wording drift a
# little without breaking the Lambda.
HEADER_ALIASES: dict[str, str] = {
    "last active date": "last_active_date",
    "lastactivedate": "last_active_date",
    "last activity date": "last_active_date",
    "acf2": "acf2",
    "login": "login",
    "user login": "login",
    "name": "name",
    "full name": "name",
    "email": "email",
    "email address": "email",
    "manager": "manager",
    "manager name": "manager",
    "manager acf2": "manager_acf2",
    "title": "title",
    "job title": "title",
    "department": "department",
    "business category": "business_category",
    "company": "company",
    "employee type": "employee_type",
    "cost centre": "cost_centre",
    "cost center": "cost_centre",
}

REQUIRED_FIELDS = ("login", "name", "email", "manager")


def build_header_index(header_row: Iterable[Any]) -> dict[str, int]:
    """Map internal field name -> column index. First occurrence wins."""
    index: dict[str, int] = {}
    seen_raw: set[str] = set()
    for position, cell in enumerate(header_row):
        key = canon(cell)
        if not key:
            continue
        if key in seen_raw:
            LOGGER.warning("Duplicate header %r at column %d ignored", key, position + 1)
            continue
        seen_raw.add(key)
        internal = HEADER_ALIASES.get(key, key.replace(" ", "_"))
        if internal in index:
            LOGGER.warning("Header %r maps to already-used field %r; ignoring", key, internal)
            continue
        index[internal] = position
    return index


# --------------------------------------------------------------------------- #
# Cell helpers
# --------------------------------------------------------------------------- #

def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = unicodedata.normalize("NFKC", str(value)).replace(" ", " ")
    return re.sub(r"\s+", " ", text).strip()


def parse_date(value: Any) -> dt.date | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial date (1900 system, with the well-known leap-year offset).
        try:
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(value))).date()
        except (OverflowError, ValueError):
            return None
    text = clean_text(value)
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y",
        "%d-%b-%Y", "%d-%b-%y", "%b %d, %Y", "%d %B %Y",
        "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M", "%d/%m/%Y %H:%M",
    ):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    LOGGER.debug("Unparseable date %r", text)
    return None


def mask_email(email: str) -> str:
    """Keep logs useful without spraying full addresses into CloudWatch."""
    if "@" not in email:
        return "***"
    local, _, domain = email.partition("@")
    head = local[:2] if len(local) > 2 else local[:1]
    return f"{head}***@{domain}"


# --------------------------------------------------------------------------- #
# Field derivation
# --------------------------------------------------------------------------- #

def resolve_market(row: dict[str, Any]) -> str | None:
    """
    Derive the market code. None means 'could not resolve'.

    The consolidated table has no market column, so the value comes from
    MARKET_SOURCE_COLUMN (default: Business Category) translated through
    MARKET_MAP, e.g. {"Hong Kong IT": "HKIT", "Philippines IT": "PHIT"}.
    """
    source_field = HEADER_ALIASES.get(CFG.market_source_column, CFG.market_source_column)
    source_field = source_field.replace(" ", "_")
    raw = clean_text(row.get(source_field))

    if CFG.market_map:
        for key, mapped in CFG.market_map.items():
            if canon(key) == canon(raw):
                return str(mapped)

    if raw and not CFG.market_map:
        # No mapping configured: pass the source value through untouched.
        return raw.upper() if CFG.market_uppercase else raw

    if CFG.market_default:
        return CFG.market_default
    return None


def resolve_active(row: dict[str, Any], today: dt.date) -> bool:
    if CFG.active_mode == "always_true":
        return True

    if CFG.active_mode == "employee_type":
        emp_type = clean_text(row.get("employee_type")).lower()
        return emp_type not in CFG.inactive_employee_types

    # default: last_active_window
    last_active = parse_date(row.get("last_active_date"))
    if last_active is None:
        return CFG.active_blank_date
    return (today - last_active).days <= CFG.active_window_days


# --------------------------------------------------------------------------- #
# Core transform  (pure -- no AWS, unit-testable)
# --------------------------------------------------------------------------- #

@dataclass
class TransformResult:
    records: list[dict[str, Any]]
    rejected: list[dict[str, Any]]
    total_rows: int
    header_index: dict[str, int]


def transform_rows(
    header_row: Iterable[Any],
    data_rows: Iterator[Iterable[Any]],
    today: dt.date | None = None,
) -> TransformResult:
    today = today or dt.datetime.now(dt.timezone.utc).date()

    index = build_header_index(header_row)
    missing = [f for f in REQUIRED_FIELDS if f not in index]
    if missing:
        raise ValueError(
            f"Consolidated sheet is missing required column(s): {', '.join(missing)}. "
            f"Columns found: {sorted(index)}"
        )

    records: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    by_login: dict[str, int] = {}
    total = 0

    for offset, raw_row in enumerate(data_rows):
        cells = list(raw_row)
        if all(c is None or clean_text(c) == "" for c in cells):
            continue  # blank spacer row
        total += 1
        excel_row = offset + 2  # header is row 1

        row = {
            fieldname: (cells[pos] if pos < len(cells) else None)
            for fieldname, pos in index.items()
        }

        def reject(reason: str) -> None:
            rejected.append({
                "row": excel_row,
                "reason": reason,
                "login": clean_text(row.get("login")),
                "email": mask_email(clean_text(row.get("email")).lower()),
            })

        email = clean_text(row.get("email")).lower()
        login = clean_text(row.get("login"))

        if not login and CFG.derive_login_from_email and email:
            login = email.split("@", 1)[0]

        if not login:
            reject("missing login")
            continue

        name = clean_text(row.get("name"))
        if not name:
            reject("missing name")
            continue

        if CFG.require_email:
            if not email:
                reject("missing email")
                continue
            if not EMAIL_RE.match(email):
                reject("malformed email")
                continue

        market = resolve_market(row)
        if market is None:
            if CFG.market_strict:
                reject(f"unmapped market (source column: {CFG.market_source_column})")
                continue
            market = ""

        record = {
            "login": login,
            "name": name,
            "email": email,
            "manager": clean_text(row.get("manager")),
            "market": market,
            "marketTag": market,
            "active": resolve_active(row, today),
            "tokens": CFG.default_tokens,
        }

        dedupe_key = login.casefold()
        if dedupe_key in by_login:
            first = records[by_login[dedupe_key]]
            if first["email"] != record["email"]:
                # Same login, different person. Silently keeping one would be
                # worse than saying so -- login is the downstream join key.
                reject(f"duplicate login collides with row for {mask_email(first['email'])}")
            else:
                reject("duplicate login (identical email) -- first occurrence kept")
            continue

        by_login[dedupe_key] = len(records)
        records.append(record)

    if CFG.sort_by_login:
        records.sort(key=lambda r: r["login"].casefold())

    return TransformResult(records, rejected, total, index)


# --------------------------------------------------------------------------- #
# Workbook reading
# --------------------------------------------------------------------------- #

def read_workbook(stream, sheet_name: str = "") -> TransformResult:
    workbook = load_workbook(stream, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Sheet {sheet_name!r} not found. Sheets: {workbook.sheetnames}"
                )
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]

        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise ValueError("Worksheet is empty -- no header row")
        return transform_rows(header_row, rows)
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# Rendering
# --------------------------------------------------------------------------- #

def render_output(records: list[dict[str, Any]]) -> str:
    body = json.dumps(records, indent=2, ensure_ascii=False)
    if CFG.output_format == "json":
        return body + "\n"
    return f"const {CFG.js_const_name} = {body};\n"


# --------------------------------------------------------------------------- #
# S3 plumbing
# --------------------------------------------------------------------------- #

def extract_s3_targets(event: dict) -> list[tuple[str, list[str]]]:
    """
    Support EventBridge 'Object Created' and raw S3 notifications alike.

    The two sources encode the key differently: S3 notifications URL-encode it
    (spaces become '+'), EventBridge generally does not. Guessing wrong turns a
    filename like 'Consolidated Users.xlsx' into a 404, so each target carries
    candidate spellings and resolve_key() picks the one that exists.
    """
    targets: list[tuple[str, list[str]]] = []

    def candidates(key: str, decoded_first: bool) -> list[str]:
        decoded = unquote_plus(key)
        order = [decoded, key] if decoded_first else [key, decoded]
        return list(dict.fromkeys(order))  # de-dupe, keep order

    detail = event.get("detail")
    if isinstance(detail, dict) and "bucket" in detail and "object" in detail:
        bucket = detail["bucket"].get("name")
        key = detail["object"].get("key")
        if bucket and key:
            targets.append((bucket, candidates(key, decoded_first=False)))

    for record in event.get("Records", []) or []:
        s3_block = record.get("s3") or {}
        bucket = (s3_block.get("bucket") or {}).get("name")
        key = (s3_block.get("object") or {}).get("key")
        if bucket and key:
            targets.append((bucket, candidates(key, decoded_first=True)))

    return targets


def resolve_key(bucket: str, candidates: list[str]) -> str:
    """Return the first candidate spelling that actually exists in the bucket."""
    last_error: Exception | None = None
    for candidate in candidates:
        try:
            s3.head_object(Bucket=bucket, Key=candidate)
            return candidate
        except ClientError as exc:
            if exc.response.get("Error", {}).get("Code") in {"404", "NoSuchKey", "NotFound"}:
                last_error = exc
                continue
            raise
    raise FileNotFoundError(
        f"None of the candidate keys exist in s3://{bucket}: {candidates}"
    ) from last_error


def put_object(bucket: str, key: str, body: str, content_type: str) -> None:
    kwargs: dict[str, Any] = {
        "Bucket": bucket,
        "Key": key,
        "Body": body.encode("utf-8"),
        "ContentType": content_type,
        "CacheControl": "no-cache, max-age=0",
    }
    if CFG.sse:
        kwargs["ServerSideEncryption"] = CFG.sse
        if CFG.sse == "aws:kms" and CFG.sse_kms_key_id:
            kwargs["SSEKMSKeyId"] = CFG.sse_kms_key_id
    s3.put_object(**kwargs)


def process_object(bucket: str, key: str) -> dict[str, Any]:
    LOGGER.info("Processing s3://%s/%s", bucket, key)

    response = s3.get_object(Bucket=bucket, Key=key)
    with response["Body"] as body:
        payload = body.read()

    result = read_workbook(io.BytesIO(payload), CFG.sheet_name)

    reject_ratio = (len(result.rejected) / result.total_rows) if result.total_rows else 0.0
    if result.total_rows and reject_ratio > CFG.max_reject_ratio:
        # A mostly-rejected file usually means the BA pasted into the wrong
        # columns. Publishing that over a good data.json is the bad outcome.
        raise ValueError(
            f"Rejected {len(result.rejected)}/{result.total_rows} rows "
            f"({reject_ratio:.0%}) -- above MAX_REJECT_RATIO "
            f"({CFG.max_reject_ratio:.0%}). Output not written. "
            f"First issues: {result.rejected[:5]}"
        )

    out_bucket = CFG.output_bucket or bucket
    put_object(out_bucket, CFG.output_key, render_output(result.records),
               "application/javascript" if CFG.output_format == "js_const" else "application/json")

    summary = {
        "source": f"s3://{bucket}/{key}",
        "output": f"s3://{out_bucket}/{CFG.output_key}",
        "generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "rowsRead": result.total_rows,
        "recordsWritten": len(result.records),
        "rowsRejected": len(result.rejected),
        "activeCount": sum(1 for r in result.records if r["active"]),
        "markets": sorted({r["market"] for r in result.records if r["market"]}),
        "rejected": result.rejected,
    }

    if CFG.report_key:
        put_object(out_bucket, CFG.report_key,
                   json.dumps(summary, indent=2, ensure_ascii=False) + "\n",
                   "application/json")

    LOGGER.info(
        "Wrote %d records (%d active, %d rejected) to s3://%s/%s",
        len(result.records), summary["activeCount"], len(result.rejected),
        out_bucket, CFG.output_key,
    )
    return summary


# --------------------------------------------------------------------------- #
# Handler
# --------------------------------------------------------------------------- #

def lambda_handler(event: dict, context: Any = None) -> dict[str, Any]:
    targets = extract_s3_targets(event)
    if not targets:
        LOGGER.warning("No S3 object in event; nothing to do. Event keys: %s", list(event))
        return {"statusCode": 400, "processed": [], "message": "no S3 object in event"}

    output_prefix = CFG.output_key.rsplit("/", 1)[0] + "/" if "/" in CFG.output_key else ""
    processed: list[dict[str, Any]] = []

    for bucket, candidates in targets:
        key = candidates[0]

        # Guard against the Lambda re-triggering on its own output.
        if output_prefix and key.startswith(output_prefix):
            LOGGER.info("Skipping own output object %s", key)
            continue
        if CFG.source_prefix and not key.startswith(CFG.source_prefix):
            LOGGER.info("Skipping %s -- outside SOURCE_PREFIX %s", key, CFG.source_prefix)
            continue
        if not key.lower().endswith((".xlsx", ".xlsm")):
            LOGGER.info("Skipping %s -- not an .xlsx/.xlsm file", key)
            continue

        try:
            processed.append(process_object(bucket, resolve_key(bucket, candidates)))
        except ClientError as exc:
            LOGGER.exception("S3 error on s3://%s/%s", bucket, key)
            raise
        except ValueError as exc:
            # Data problems are the operator's to fix, not a retryable fault --
            # but they must be loud, so still raise after logging clearly.
            LOGGER.error("Validation failure on s3://%s/%s: %s", bucket, key, exc)
            raise

    return {"statusCode": 200, "processed": processed}